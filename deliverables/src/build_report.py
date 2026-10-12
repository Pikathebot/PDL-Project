# -*- coding: utf-8 -*-
"""Build the Sentinel PDL-2 report (.docx + .pdf) in the R.M.K. sample-report format.

    python build_report.py            # two passes: lay out, read page numbers from Word's PDF, rebuild

Requires python-docx, pywin32 (Word for PDF export) and PyMuPDF (page-number pass).
"""
import os, re, sys
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)                 # deliverables/
REPO = os.path.dirname(ROOT)                 # repo root (code excerpts)
sys.path.insert(0, HERE)
from report_content import *                 # noqa
from report_content2 import BODY2, REFERENCES

OUT_DOCX = os.path.join(ROOT, "Sentinel_PDL2_Report.docx")
OUT_PDF = os.path.join(ROOT, "Sentinel_PDL2_Report.pdf")
FONT = "Times New Roman"
MAX_IMG_H = 7.0   # inches; keeps a figure + caption on one page

# ---------------------------------------------------------------- helpers
def set_cell_border_none(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for b in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{b}"); el.set(qn("w:val"), "nil"); borders.append(el)
    tblPr.append(borders)

def page_number_field(paragraph):
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind); run._r.append(el)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text; run._r.append(el)
    run.font.name = FONT; run.font.size = Pt(11)

def set_page_numbering(section, fmt=None, start=None):
    sectPr = section._sectPr
    pg = sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType"); sectPr.append(pg)
    if fmt: pg.set(qn("w:fmt"), fmt)
    if start is not None: pg.set(qn("w:start"), str(start))

def setup_section(section, numbered):
    section.page_height, section.page_width = Inches(11.69), Inches(8.27)
    section.top_margin = section.bottom_margin = Inches(1.0)
    section.left_margin, section.right_margin = Inches(1.25), Inches(1.0)
    section.footer.is_linked_to_previous = False
    fp = section.footer.paragraphs[0]
    for r in list(fp.runs): r._r.getparent().remove(r._r)
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if numbered:
        page_number_field(fp)

class R:
    def __init__(self):
        self.doc = Document()
        st = self.doc.styles["Normal"]
        st.font.name = FONT; st.font.size = Pt(12)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        pf = st.paragraph_format
        pf.line_spacing = 1.5; pf.space_after = Pt(6); pf.space_before = Pt(0)
        for name in ("List Bullet", "List Bullet 2"):
            s = self.doc.styles[name]; s.font.name = FONT; s.font.size = Pt(12)
        setup_section(self.doc.sections[0], numbered=False)

    # -- primitives
    def para(self, text="", bold=False, italic=False, size=12, align="justify", after=6, before=0,
             spacing=1.5, keep=False, indent=None):
        p = self.doc.add_paragraph()
        p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                       "left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT}[align]
        pf = p.paragraph_format
        pf.space_after, pf.space_before, pf.line_spacing = Pt(after), Pt(before), spacing
        pf.keep_with_next = keep
        if indent is not None: pf.left_indent = Inches(indent)
        self.runs(p, text, bold, italic, size)
        return p

    def runs(self, p, text, bold=False, italic=False, size=12):
        # **bold** spans
        for i, part in enumerate(re.split(r"\*\*", text)):
            if not part: continue
            r = p.add_run(part); r.font.name = FONT; r.font.size = Pt(size)
            r.bold = bold or (i % 2 == 1); r.italic = italic

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def new_section(self, fmt, start):
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        setup_section(sec, numbered=True)
        set_page_numbering(sec, fmt, start)

    def image(self, path, width, height_cap=MAX_IMG_H):
        w, h = Image.open(path).size
        width = min(width, 6.0)
        if width * h / w > height_cap:
            width = height_cap * w / h
        p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.0
        p.add_run().add_picture(path, width=Inches(width))

    def bullets(self, items, level=1):
        style = "List Bullet" if level == 1 else "List Bullet 2"
        for it in items:
            p = self.doc.add_paragraph(style=style)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.5
            m = re.match(r"^([^–:]{2,45}?)( – |: )(.*)$", it)
            if m:
                self.runs(p, m.group(1), bold=True); self.runs(p, m.group(2) + m.group(3))
            else:
                self.runs(p, it)

    def table(self, headers, rows):
        t = self.doc.add_table(rows=1, cols=len(headers)); t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(headers):
            c = t.rows[0].cells[i]; c.text = ""
            self.runs(c.paragraphs[0], h, bold=True); c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""; self.runs(cells[i].paragraphs[0], v)
        for row in t.rows:
            row.cells[0].width = Inches(2.0); row.cells[1].width = Inches(4.0)
            for c in row.cells:
                for p in c.paragraphs:
                    p.paragraph_format.line_spacing = 1.15; p.paragraph_format.space_after = Pt(3)
        self.para("", after=2)

    def code(self, title, text):
        self.para(title, bold=True, align="left", after=4, before=8, keep=True)
        for line in text.rstrip("\n").split("\n"):
            p = self.doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf = p.paragraph_format; pf.space_after = Pt(0); pf.line_spacing = 1.0
            r = p.add_run(line.replace("\t", "    ") or " ")
            r.font.name = "Courier New"; r.font.size = Pt(8.5)
            r._element.rPr.rFonts.set(qn("w:hAnsi"), "Courier New")

# ---------------------------------------------------------------- pages
def cover(r):
    sp = 1.15
    r.para(COURSE, bold=True, size=14, align="center", after=18, spacing=sp)
    for line in TITLE_LINES:
        r.para(line, bold=True, size=14, align="center", after=2, spacing=sp)
    r.para("", after=10)
    r.para("A PROJECT REPORT", bold=True, align="center", after=6, spacing=sp)
    r.para("Submitted by", bold=True, italic=True, align="center", after=6, spacing=sp)
    t = r.doc.add_table(rows=0, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER; set_cell_border_none(t)
    for name, reg in STUDENTS:
        c = t.add_row().cells
        for cell, txt, al in ((c[0], name, WD_ALIGN_PARAGRAPH.LEFT), (c[1], f"({reg})", WD_ALIGN_PARAGRAPH.LEFT)):
            cell.text = ""; r.runs(cell.paragraphs[0], txt, bold=True)
            cell.paragraphs[0].alignment = al; cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            cell.paragraphs[0].paragraph_format.line_spacing = 1.15
        c[0].width = Inches(2.6); c[1].width = Inches(1.6)
    r.para("", after=4)
    r.para("in partial fulfillment for the award of the degree of", bold=True, italic=True, align="center", after=6, spacing=sp)
    r.para("BACHELOR OF ENGINEERING", bold=True, size=13, align="center", after=4, spacing=sp)
    r.para("IN", bold=True, align="center", after=4, spacing=sp)
    r.para("COMPUTER SCIENCE AND ENGINEERING", bold=True, align="center", after=8, spacing=sp)
    r.para("R.M.K. ENGINEERING COLLEGE", bold=True, align="center", after=0, spacing=sp)
    r.para("(An Autonomous Institution)", align="center", after=0, spacing=sp)
    r.para("R.S.M. Nagar, Kavaraipettai-601 206", bold=True, size=11, align="center", after=8, spacing=sp)
    r.image(os.path.join(ROOT, "assets/misc/rmk_logo.png"), 1.25)
    r.para("", after=4)
    r.para(MONTH, bold=True, align="center", after=0, spacing=sp)

def bonafide(r):
    r.page_break()
    r.para("R.M.K. ENGINEERING COLLEGE", bold=True, size=14, align="center", after=0, spacing=1.15)
    r.para("(An Autonomous Institution)", bold=True, size=11, align="center", after=0, spacing=1.15)
    r.para("R.S.M. Nagar, Kavaraipettai-601 206", bold=True, size=11, align="center", after=24, spacing=1.15)
    r.para("BONAFIDE CERTIFICATE", bold=True, size=14, align="center", after=18)
    r.para(f"Certified that this Product Development Lab – II report **“{TITLE}”** is the Bonafide work of "
           f"**{STUDENTS_INLINE}** who carried out the **{COURSE.replace('LAB', 'Lab').replace('PRODUCT DEVELOPMENT', 'Product Development')}** "
           "work under my supervision.", after=36)
    t = r.doc.add_table(rows=1, cols=2); set_cell_border_none(t)
    for cell, lines in ((t.rows[0].cells[0], HOD), (t.rows[0].cells[1], SUPERVISOR)):
        cell.text = ""
        p = cell.paragraphs[0]; r.runs(p, "SIGNATURE", bold=True); p.paragraph_format.space_after = Pt(30)
        for i, line in enumerate(lines):
            q = cell.add_paragraph(); r.runs(q, line, bold=i < 2)
            q.paragraph_format.space_after = Pt(0); q.paragraph_format.line_spacing = 1.15
    r.para("", after=30)
    r.para("Submitted for the Product Development Lab – 2 Viva-Voce held on ………………………… at "
           "R.M.K. Engineering College, Kavaraipettai, Tiruvallur District– 601206.", after=60, spacing=1.5)
    r.para("INTERNAL EXAMINER", bold=True, align="left", after=0)

ACK = [
    "We earnestly portray our sincere gratitude and regard to our beloved **Chairman Shri. R. S. Munirathinam, our "
    "Vice Chairman, Shri. R. M. Kishore** and our **Director, Shri. R. Jyothi Naidu,** for the interest and affection "
    "shown towards us throughout the course.",
    "We convey our sincere thanks to our **Principal, Dr. K. A. Mohamed Junaid,** for being the source of inspiration "
    "in this college.",
    "We reveal our sincere thanks to our **Professor and Head of the Department, Computer Science and Engineering, "
    "Dr. T. Sethukarasi,** for her commendable support and encouragement for the completion of our project.",
    "We would like to express our sincere gratitude to our **Product Development Lab Incharge and Supervisor "
    "Ms. Sumitha T, Assistant Professor,** for the valuable suggestions and guidance towards the successful "
    "completion of this project in a global manner.",
    "We take this opportunity to extend our thanks to all faculty members of the Department of Computer Science and "
    "Engineering, Parents and friends for all that they meant to us during the crucial times of the completion of "
    "our project.",
]

def front_matter(r, pages):
    r.new_section("lowerRoman", 3)
    r.para("ACKNOWLEDGEMENT", bold=True, size=14, align="center", after=18)
    for a in ACK:
        r.para(a, after=14, spacing=2.0)
    r.page_break()
    r.para("ABSTRACT", bold=True, size=14, align="center", after=18)
    r.para(ABSTRACT, spacing=2.0)
    r.page_break()
    # Table of contents
    r.para("TABLE OF CONTENTS", bold=True, size=14, align="center", after=12)
    rows = [("", "ABSTRACT", pages.get("ABSTRACT", "iv"), True),
            ("", "LIST OF FIGURES", pages.get("LIST OF FIGURES", "vii"), True),
            ("", "LIST OF ABBREVIATIONS", pages.get("LIST OF ABBREVIATIONS", "viii"), True)]
    for blk in BODY + BODY2:
        if blk[0] == "chapter":
            rows.append((str(blk[1]), blk[2], pages.get(f"CH{blk[1]}", "00"), True))
        elif blk[0] in ("h2",) or (blk[0] == "h3" and re.match(r"^\d+\.\d+\.\d+ ", blk[1])):
            rows.append(("", blk[1], pages.get(blk[1], "00"), False))
    rows.append(("", "REFERENCES", pages.get("REFERENCES", "00"), True))
    t = r.doc.add_table(rows=1, cols=3); set_cell_border_none(t)
    for i, h in enumerate(("CHAPTER", "TITLE", "PAGE NO")):
        c = t.rows[0].cells[i]; c.text = ""; r.runs(c.paragraphs[0], h, bold=True)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for ch, title, pg, bold in rows:
        cells = t.add_row().cells
        indent = 0 if bold else (0.3 if title.count(".") < 2 or re.match(r"^\d+\.\d+ ", title) else 0.6)
        for i, (txt, al) in enumerate(((ch, WD_ALIGN_PARAGRAPH.CENTER), (title, WD_ALIGN_PARAGRAPH.LEFT),
                                         (pg, WD_ALIGN_PARAGRAPH.CENTER))):
            cells[i].text = ""
            p = cells[i].paragraphs[0]; p.alignment = al
            p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.15
            if i == 1:
                p.paragraph_format.left_indent = Inches(indent)
                txt = re.sub(r"^(\d+\.\d+\.\d+ )", r"\1", txt)
                r.runs(p, txt if bold else (txt if re.match(r"^\d+\.\d+\.\d+ ", txt) else txt.upper()), bold=bold, size=11)
            else:
                r.runs(p, txt, bold=bold and i == 0, size=11)
    for row in t.rows:
        row.cells[0].width, row.cells[1].width, row.cells[2].width = Inches(0.9), Inches(4.1), Inches(1.1)
    r.page_break()
    # List of figures
    r.para("LIST OF FIGURES", bold=True, size=14, align="center", after=12)
    t = r.doc.add_table(rows=1, cols=3); set_cell_border_none(t)
    for i, h in enumerate(("Figure No", "Name", "Page No")):
        c = t.rows[0].cells[i]; c.text = ""; r.runs(c.paragraphs[0], h, bold=True)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    for num, name in FIGURES:
        cells = t.add_row().cells
        for i, (txt, al) in enumerate(((num, WD_ALIGN_PARAGRAPH.CENTER), (name, WD_ALIGN_PARAGRAPH.LEFT),
                                       (pages.get("FIG" + num, "00"), WD_ALIGN_PARAGRAPH.CENTER))):
            cells[i].text = ""; p = cells[i].paragraphs[0]; p.alignment = al
            p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.3
            r.runs(p, txt, size=11)
    for row in t.rows:
        row.cells[0].width, row.cells[1].width, row.cells[2].width = Inches(1.2), Inches(3.9), Inches(1.0)
    r.page_break()
    r.para("LIST OF ABBREVIATIONS", bold=True, size=14, align="center", after=12)
    t = r.doc.add_table(rows=0, cols=2); set_cell_border_none(t)
    for a, full in ABBREVIATIONS:
        cells = t.add_row().cells
        for i, txt in enumerate((a, full)):
            cells[i].text = ""; p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.15
            r.runs(p, txt, bold=(i == 0))
    for row in t.rows:
        row.cells[0].width, row.cells[1].width = Inches(1.3), Inches(4.7)

def body(r):
    r.new_section("decimal", 1)
    first = True
    for blk in BODY + BODY2:
        k = blk[0]
        if k == "chapter":
            h = r.para(f"CHAPTER {blk[1]}", bold=True, size=14, align="center", after=6, keep=True)
            h.paragraph_format.page_break_before = not first
            first = False
            r.para(blk[2], bold=True, size=14, align="center", after=14, keep=True)
        elif k == "h2":
            r.para(blk[1], bold=True, size=13, align="left", before=8, after=6, keep=True)
        elif k == "h3":
            r.para(blk[1], bold=True, size=12, align="left", before=6, after=4, keep=True)
        elif k == "p":
            r.para(blk[1])
        elif k == "bul":
            r.bullets(blk[1])
        elif k == "sub":
            r.bullets(blk[1], level=2)
        elif k == "fig":
            _, f, num, cap, w = blk
            r.image(os.path.join(ROOT, f), w)
            r.para(f"Fig {num} {cap}", size=11, align="center", after=14, spacing=1.0)
        elif k == "table":
            r.table(blk[1], blk[2])
        elif k == "codefile":
            _, title, path, a, b = blk
            lines = open(os.path.join(REPO, path), encoding="utf-8").read().split("\n")[a - 1:b]
            r.code(title, "\n".join(lines))
    r.para("REFERENCES", bold=True, size=14, align="left", after=12).paragraph_format.page_break_before = True
    for i, ref in enumerate(REFERENCES, 1):
        p = r.para(f"[{i}]  {ref}", after=6, spacing=1.15)
        p.paragraph_format.left_indent = Inches(0.45); p.paragraph_format.first_line_indent = Inches(-0.45)

def build(pages):
    r = R()
    cover(r); bonafide(r); front_matter(r, pages); body(r)
    r.doc.save(OUT_DOCX)

# ---------------------------------------------------------------- page numbers
def export_pdf():
    import win32com.client
    word = win32com.client.DispatchEx("Word.Application"); word.Visible = False
    try:
        d = word.Documents.Open(OUT_DOCX, ReadOnly=True)
        d.Repaginate()
        d.SaveAs2(OUT_PDF, FileFormat=17); d.Close(False)
    finally:
        word.Quit()

def roman(n):
    vals = [(10, "x"), (9, "ix"), (5, "v"), (4, "iv"), (1, "i")]
    out = ""
    for v, s in vals:
        while n >= v: out += s; n -= v
    return out

def norm(s):
    return re.sub(r"\s+", " ", s).strip().lower()

def read_pages():
    import pymupdf
    doc = pymupdf.open(OUT_PDF)
    # Drop the footer page number, which PyMuPDF extracts first.
    texts = [re.sub(r"^([ivxl]+|\d+) ", "", norm(p.get_text())) for p in doc]
    pages = {}
    def find(needle, start=0):
        n = norm(needle)
        for i in range(start, len(texts)):
            if n in texts[i]: return i
        return None
    # front matter: page index 2 is "iii"
    for key in ("ABSTRACT", "LIST OF FIGURES", "LIST OF ABBREVIATIONS"):
        i = find(key if key != "ABSTRACT" else "abstract emergency response", 2)
        if key == "ABSTRACT":
            i = next(j for j in range(2, 8) if texts[j].startswith("abstract"))
        if key == "LIST OF FIGURES":
            i = next(j for j in range(4, 10) if texts[j].startswith("list of figures"))
        if key == "LIST OF ABBREVIATIONS":
            i = next(j for j in range(4, 12) if texts[j].startswith("list of abbreviations"))
        pages[key] = roman(i + 1)
    body0 = next(j for j in range(len(texts)) if texts[j].startswith("chapter 1 introduction")
                 or re.match(r"^chapter 1 introduction", texts[j]) or "chapter 1 introduction overview" in texts[j][:60])
    cur = body0
    for blk in BODY + BODY2:
        if blk[0] == "chapter":
            i = find(f"chapter {blk[1]} {blk[2]}", cur); cur = i
            pages[f"CH{blk[1]}"] = f"{i - body0 + 1:02d}"
        elif blk[0] in ("h2", "h3"):
            i = find(blk[1], cur)
            if i is not None:
                cur = i; pages[blk[1]] = f"{i - body0 + 1:02d}"
        elif blk[0] == "fig":
            i = find(f"fig {blk[2]} {blk[3]}", cur)
            if i is not None:
                cur = i; pages["FIG" + blk[2]] = f"{i - body0 + 1:02d}"
    pages["REFERENCES"] = f"{find('references [1]', cur) - body0 + 1:02d}"
    print(f"PDF pages: {len(texts)}  (front {body0}, body {len(texts) - body0})")
    return pages

if __name__ == "__main__":
    build({}); export_pdf(); p1 = read_pages()
    build(p1); export_pdf(); p2 = read_pages()
    if p1 != p2:
        print("pagination shifted; third pass"); build(p2); export_pdf(); read_pages()
    print("wrote", OUT_DOCX, "and", OUT_PDF)
